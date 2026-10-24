"""アポツールのシフト表CSVを、月ごとのシフトカレンダー(Excel)に変換する。

使い方:
    python3 make_calendar.py <シフト表.csv> [出力先.xlsx]

- CSVはアポツールが書き出す Shift-JIS のもの(UTF-8 でも可)
- 記号: フ=1日出勤(名前のみ) / A=午前(名前AM) / P=午後(名前PM) / 休=記入しない
- 1マスは職種ごとに改行(歯科医 / 衛生士 / 助手 の順。CSVの見出し順に従う)
- 全員が休みの日は「休診」
- 名前の先頭の DR/DH/DA は外す。「担当希望なし」の行は使わない
- 書き出したあと、Excelを読み戻してCSVと1人×1日ずつ照合し、食い違いがあれば失敗する
"""
import calendar
import copy
import csv
import datetime
import io
import re
import sys
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font

TEMPLATE = Path(__file__).with_name("template.xlsx")
COLS = "BDFHJLN"          # 月〜日の日付列(右隣の列と結合して名前欄になる)
FIRST_ROW = 5             # 1週目の日付行。名前行はその1つ下
MAX_WEEKS = 6
NAME_ROW_HEIGHT = 85
SKIP_ROWS = {"担当希望なし"}
PREFIXES = ("DR", "DH", "DA")
CODES = {"フ": "", "A": "AM", "P": "PM", "休": None}

# 日付の文字色(テンプレートの元の配色)
BLACK, BLUE, RED = None, "FF3366FF", "FFFF0000"
GRAY, GRAY_BLUE, GRAY_RED = "FFC0C0C0", "FFC1D1FF", "FFFFB2B2"


# ---------------------------------------------------------------- 祝日
def _nth_monday(y, m, n):
    d = datetime.date(y, m, 1)
    d += datetime.timedelta(days=(7 - d.weekday()) % 7)
    return d + datetime.timedelta(weeks=n - 1)


def japanese_holidays(year):
    """{date: 祝日名}。春分・秋分は 1980〜2099 年で有効な近似式。"""
    y = year
    h = {
        datetime.date(y, 1, 1): "元日",
        _nth_monday(y, 1, 2): "成人の日",
        datetime.date(y, 2, 11): "建国記念の日",
        datetime.date(y, 2, 23): "天皇誕生日",
        datetime.date(y, 3, int(20.8431 + 0.242194 * (y - 1980) - (y - 1980) // 4)): "春分の日",
        datetime.date(y, 4, 29): "昭和の日",
        datetime.date(y, 5, 3): "憲法記念日",
        datetime.date(y, 5, 4): "みどりの日",
        datetime.date(y, 5, 5): "こどもの日",
        _nth_monday(y, 7, 3): "海の日",
        datetime.date(y, 8, 11): "山の日",
        _nth_monday(y, 9, 3): "敬老の日",
        datetime.date(y, 9, int(23.2488 + 0.242194 * (y - 1980) - (y - 1980) // 4)): "秋分の日",
        _nth_monday(y, 10, 2): "スポーツの日",
        datetime.date(y, 11, 3): "文化の日",
        datetime.date(y, 11, 23): "勤労感謝の日",
    }
    # 国民の休日: 祝日にはさまれた平日
    for d in sorted(h):
        mid, nxt = d + datetime.timedelta(1), d + datetime.timedelta(2)
        if nxt in h and mid not in h and mid.weekday() != 6:
            h[mid] = "国民の休日"
    # 振替休日: 日曜の祝日の後の最初の平日
    for d in sorted(h):
        if d.weekday() == 6 and h[d] != "振替休日":
            s = d + datetime.timedelta(1)
            while s in h:
                s += datetime.timedelta(1)
            h[s] = "振替休日"
    return h


def holidays_around(year):
    h = {}
    for y in (year - 1, year, year + 1):
        h.update(japanese_holidays(y))
    return h


# ---------------------------------------------------------------- CSV
def read_csv(path):
    raw = Path(path).read_bytes()
    for enc in ("utf-8-sig", "cp932"):
        try:
            text = raw.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    else:
        raise SystemExit("CSVの文字コードが読めません")
    rows = list(csv.reader(io.StringIO(text)))

    m = re.search(r"(\d{4})年(\d{1,2})月", rows[0][0] if rows[0] else "")
    if not m:
        raise SystemExit(f"1行目から年月が読めません: {rows[0]}")
    year, month = int(m.group(1)), int(m.group(2))
    ndays = calendar.monthrange(year, month)[1]

    days = [int(x) for x in rows[1][1:] if x.strip()]
    if days != list(range(1, ndays + 1)):
        raise SystemExit(f"日付の行が {year}年{month}月 の日数と合いません: {days}")
    wd = "月火水木金土日"
    for i, w in enumerate(rows[2][1:ndays + 1], 1):
        if wd[datetime.date(year, month, i).weekday()] != w.strip():
            raise SystemExit(f"{month}/{i} の曜日がCSV({w})と実際({wd[datetime.date(year, month, i).weekday()]})で違います")

    sections = []          # [(見出し, [(名前, [コード...])])]
    for r in rows[3:]:
        if not r or not r[0].strip():
            continue
        head = r[0].strip()
        if head == "勤務時間帯":
            break
        if len([x for x in r if x.strip()]) == 1:
            sections.append((head, []))
            continue
        if not sections:
            raise SystemExit(f"職種の見出しより前に行があります: {head}")
        codes = [x.strip() for x in r[1:ndays + 1]]
        bad = {c for c in codes if c not in CODES}
        if bad or len(codes) != ndays:
            raise SystemExit(f"{head} の行に想定外の記号/日数があります: {bad or len(codes)}")
        if head in SKIP_ROWS:
            continue
        name = head[2:] if head[:2] in PREFIXES else head
        sections[-1][1].append((name, codes))
    return year, month, ndays, sections


def day_lines(sections, d):
    lines = []
    for _, people in sections:
        names = [n + CODES[c[d - 1]] for n, c in people if CODES[c[d - 1]] is not None]
        lines.append("・".join(names))
    return lines


# ---------------------------------------------------------------- Excel
def build(csv_path, out_path):
    year, month, ndays, sections = read_csv(csv_path)
    hol = holidays_around(year)
    wb = openpyxl.load_workbook(TEMPLATE)
    ws = wb.active
    ws.title = f"{month}月"
    ws["B2"].value = year
    ws["H2"].value = month

    first = datetime.date(year, month, 1)
    start = first - datetime.timedelta(days=first.weekday())
    weeks = (first.weekday() + ndays + 6) // 7

    for w in range(MAX_WEEKS):
        r = FIRST_ROW + w * 2
        if w >= weeks:                        # 使わない週は消す
            for rr in (r, r + 1):
                for c in COLS:
                    rng = f"{c}{rr}:{chr(ord(c) + 1)}{rr}"
                    if rng in {str(x) for x in ws.merged_cells.ranges}:
                        ws.unmerge_cells(rng)
                for col in range(2, 16):
                    ws.cell(rr, col).value = None
                    ws.cell(rr, col)._style = copy.copy(ws.cell(3, 1)._style)
                ws.row_dimensions[rr].height = None
            continue
        ws.row_dimensions[r + 1].height = NAME_ROW_HEIGHT
        for i, c in enumerate(COLS):
            d = start + datetime.timedelta(days=w * 7 + i)
            inside = d.month == month
            hname = hol.get(d)
            if d.weekday() == 6 or hname:
                color = RED if inside else GRAY_RED
            elif d.weekday() == 5:
                color = BLUE if inside else GRAY_BLUE
            else:
                color = BLACK if inside else GRAY
            date_cell = ws[f"{c}{r}"]
            date_cell.value = d.day
            date_cell.font = Font(name=date_cell.font.name, sz=16, b=True, color=color)
            label = ws.cell(r, ws[f"{c}{r}"].column + 1)
            label.value = hname
            label.font = Font(name=date_cell.font.name, sz=11, color=RED if inside else GRAY_RED)

            cell = ws[f"{c}{r + 1}"]
            cell.value = None
            if inside:
                lines = day_lines(sections, d.day)
                closed = not any(lines)
                cell.value = "休診" if closed else "\n".join(lines)
                cell.font = Font(name=date_cell.font.name, sz=36 if closed else 11)
                cell.alignment = Alignment(wrap_text=True, vertical="top",
                                           horizontal="center" if closed else "left")

    wb.save(out_path)
    verify(out_path, year, month, ndays, sections)
    return year, month


def verify(out_path, year, month, ndays, sections):
    """書き出したExcelを読み戻し、CSVと1人×1日ずつ照合する。"""
    ws = openpyxl.load_workbook(out_path).active
    first = datetime.date(year, month, 1)
    start = first - datetime.timedelta(days=first.weekday())
    errors = []
    seen = set()
    for w in range(MAX_WEEKS):
        r = FIRST_ROW + w * 2
        for i, c in enumerate(COLS):
            d = start + datetime.timedelta(days=w * 7 + i)
            v = ws[f"{c}{r}"].value
            if v is None:
                continue
            if v != d.day:
                errors.append(f"{c}{r}: 日付 {v} != {d.day}")
            text = ws[f"{c}{r + 1}"].value or ""
            if d.month != month:
                if text:
                    errors.append(f"{d}: 前後の月のマスに記入あり")
                continue
            seen.add(d.day)
            lines = ["" for _ in sections] if text == "休診" else text.split("\n")
            if len(lines) != len(sections):
                errors.append(f"{month}/{d.day}: 段の数が {len(lines)}")
                continue
            for (sec, people), line in zip(sections, lines):
                got = {}
                for tok in filter(None, line.split("・")):
                    base, code = (tok[:-2], "A") if tok.endswith("AM") else (tok[:-2], "P") if tok.endswith("PM") else (tok, "フ")
                    got[base] = code
                for name, codes in people:
                    g = got.pop(name, "休")
                    if g != codes[d.day - 1]:
                        errors.append(f"{month}/{d.day} {sec} {name}: CSV={codes[d.day - 1]} Excel={g}")
                for name in got:
                    errors.append(f"{month}/{d.day} {sec}: CSVにない名前 {name}")
    if seen != set(range(1, ndays + 1)):
        errors.append(f"カレンダーにない日: {sorted(set(range(1, ndays + 1)) - seen)}")
    if errors:
        raise SystemExit("照合エラー:\n" + "\n".join(errors))
    n = sum(len(p) for _, p in sections)
    print(f"照合OK: {n}人 × {ndays}日、食い違い0件")


if __name__ == "__main__":
    if len(sys.argv) not in (2, 3):
        raise SystemExit(__doc__)
    src = sys.argv[1]
    if len(sys.argv) == 3:
        out = sys.argv[2]
    else:
        y, m, _, _ = read_csv(src)
        out = str(Path(src).with_name(f"{y}年{m}月シフト表.xlsx"))
    build(src, out)
    print(f"作成: {out}")
